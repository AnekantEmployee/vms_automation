from pptx.enum.text import PP_ALIGN
from config.colors import HEADER_BG, HEADER_TEXT
from config.sizes import HEADER_LEFT, HEADER_TOP, HEADER_WIDTH, HEADER_HEIGHT, HEADER_FONT

def create_header(slide, text, left=HEADER_LEFT, top=HEADER_TOP, 
                 width=HEADER_WIDTH, height=HEADER_HEIGHT, 
                 font_size=HEADER_FONT, bg_color=HEADER_BG, text_color=HEADER_TEXT):
    """
    Create a styled header textbox on a slide
    """
    header = slide.shapes.add_textbox(left, top, width, height)
    
    # Style the header
    header_frame = header.text_frame
    header_frame.text = text
    header_para = header_frame.paragraphs[0]
    header_para.alignment = PP_ALIGN.LEFT
    
    # Style header text
    header_run = header_para.runs[0]
    header_run.font.size = font_size
    header_run.font.bold = True
    header_run.font.color.rgb = text_color
    
    # Add colored background to header
    header.fill.solid()
    header.fill.fore_color.rgb = bg_color
    
    return header